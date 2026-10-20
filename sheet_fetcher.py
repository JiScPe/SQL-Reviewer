import re
import random
import string
from pathlib import Path
from typing import List, Tuple, Optional
from dataclasses import dataclass

from playwright.sync_api import sync_playwright, Browser, BrowserContext, Page

from config import get_config
from database import init_db, create_run, update_run_state, add_event


@dataclass
class SQLFile:
    sr_number: str
    content: str
    filename: str


class SheetFetcher:
    def __init__(self):
        self.config = get_config()
        self.browser: Optional[Browser] = None
        self.context: Optional[BrowserContext] = None
        self.page: Optional[Page] = None
        self.sr_pattern = re.compile(
            r"(?<![A-Za-z0-9])(?:" + "|".join(self.config.sr_prefixes) + r")\d{9}(?!\d)"
        )
        self.sep_pattern = re.compile(r"\((\d{2})\.(\d{2}) online\)")
        self.full_sep_pattern = re.compile(r"(\d{4})\.(\d{2})\.(\d{2})\s+(\d{1,2}:\d{2})\s+package\s+\((\d{2})\.(\d{2}) online\)")

    def fetch(self, deploy_date: str, run_id: int) -> List[SQLFile]:
        """Fetch SQL files for the given deploy date."""
        init_db()
        
        # Validate deploy_date format
        if not re.match(r"^\d{2}\.\d{2}$", deploy_date):
            raise ValueError(f"Invalid deploy date format: {deploy_date}. Expected MM.DD")
        
        update_run_state(run_id, "Opening sheet")
        add_event(run_id, "Opening sheet", "started", f"Opening kdocs sheet for deploy {deploy_date}")
        
        self._launch_browser()
        
        try:
            self._navigate_and_wait_for_login(run_id)
            xlsx_path = self._export_to_xlsx(run_id)
            update_run_state(run_id, "Selecting deploy rows")
            add_event(run_id, "Selecting deploy rows", "started", "Parsing deploy window from exported sheet")
            
            sql_files = self._parse_and_download(xlsx_path, deploy_date, run_id)
            
            add_event(run_id, "Selecting deploy rows", "completed", f"Found {len(sql_files)} SQL file(s)")
            return sql_files
        finally:
            self._close_browser()

    def _launch_browser(self):
        """Launch Playwright browser with persistent profile."""
        self.browser = sync_playwright().start().chromium.launch_persistent_context(
            user_data_dir=self.config.browser_profile_dir,
            headless=False,
            args=["--disable-blink-features=AutomationControlled"]
        )
        self.context = self.browser
        self.page = self.context.new_page()

    def _close_browser(self):
        if self.context:
            self.context.close()
        if self.browser:
            # Note: persistent context doesn't have a separate browser to close
            pass

    def _navigate_and_wait_for_login(self, run_id: int):
        """Navigate to kdocs and wait for login if needed."""
        self.page.goto(self.config.kdocs_sheet_url)
        add_event(run_id, "Opening sheet", "progress", f"Navigated to {self.config.kdocs_base_url}")
        
        # Wait for either login page or sheet to be visible
        login_detected = False
        for _ in range(self.config.login_timeout_seconds * 2):  # check every 0.5s
            # Check for login QR code
            if self.page.locator('text="微信扫码"').count() > 0 or \
               self.page.locator('text="WeChat"').count() > 0 or \
               self.page.locator('[class*="qr"]').count() > 0:
                if not login_detected:
                    update_run_state(run_id, "Waiting for login")
                    add_event(run_id, "Waiting for login", "started", "WeChat QR code detected - waiting for user to scan")
                    login_detected = True
            
            # Check if sheet is loaded (look for spreadsheet grid)
            if self.page.locator('[class*="spreadsheet"]').count() > 0 or \
               self.page.locator('[class*="grid"]').count() > 0:
                if login_detected:
                    add_event(run_id, "Waiting for login", "completed", "Login successful, sheet loaded")
                break
            
            self.page.wait_for_timeout(500)
        else:
            raise TimeoutError(f"Login timeout after {self.config.login_timeout_seconds} seconds")

    def _export_to_xlsx(self, run_id: int) -> str:
        """Export the current sheet to .xlsx file."""
        add_event(run_id, "Opening sheet", "progress", "Exporting sheet to .xlsx")
        
        # Click File menu -> Download as -> Excel
        # This is a simplified implementation - actual selectors depend on kdocs UI
        self.page.locator('text="文件"').first.click()  # File menu
        self.page.wait_for_timeout(500)
        self.page.locator('text="下载为"').first.click()  # Download as
        self.page.wait_for_timeout(500)
        self.page.locator('text="Excel"').first.click()  # Excel format
        
        # Wait for download
        with self.page.expect_download() as download_info:
            pass
        download = download_info.value
        xlsx_path = f"workspace/deploy_{run_id}.xlsx"
        Path(xlsx_path).parent.mkdir(parents=True, exist_ok=True)
        download.save_as(xlsx_path)
        
        add_event(run_id, "Opening sheet", "completed", f"Sheet exported to {xlsx_path}")
        return xlsx_path

    def _parse_and_download(self, xlsx_path: str, deploy_date: str, run_id: int) -> List[SQLFile]:
        """Parse the .xlsx file and download SQL files."""
        import openpyxl
        
        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb.active
        
        # Find all separator rows
        separators = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, values_only=False), 1):
            cell_a = row[0] if row else None
            if cell_a and cell_a.value:
                text = str(cell_a.value)
                if self.sep_pattern.search(text):
                    separators.append((row_idx, text))
        
        # Find target separator
        target_sep_idx = None
        for idx, (row_num, text) in enumerate(separators):
            if deploy_date in text:
                target_sep_idx = idx
                break
        
        if target_sep_idx is None:
            raise ValueError(f"Deploy separator for {deploy_date} not found")
        
        # Determine window boundaries
        target_row = separators[target_sep_idx][0]
        if target_sep_idx > 0:
            start_row = separators[target_sep_idx - 1][0] + 1
        else:
            start_row = 2  # skip header
        
        end_row = target_row - 1
        
        if start_row > end_row:
            add_event(run_id, "Selecting deploy rows", "completed", "No SQL files in this deploy")
            return []
        
        # Extract SQL files from window
        sql_files = []
        sr_counts = {}
        
        for row_idx in range(start_row, end_row + 1):
            row = ws[row_idx]
            if len(row) < 12:  # Column L is index 11
                continue
            
            cell_l = row[11]  # Column L (0-indexed)
            if not cell_l or not cell_l.hyperlink:
                continue
            
            # Extract SR from Column D (index 3)
            cell_d = row[3] if len(row) > 3 else None
            sr_number = None
            if cell_d and cell_d.value:
                text = str(cell_d.value)
                match = self.sr_pattern.search(text)
                if match:
                    sr_number = match.group(0)
            
            # Generate filename
            if sr_number:
                base_name = f"SQL_{sr_number}.sql"
                count = sr_counts.get(sr_number, 0) + 1
                sr_counts[sr_number] = count
                if count > 1:
                    filename = f"SQL_{sr_number}_{count}.sql"
                else:
                    filename = base_name
            else:
                # Fallback naming
                date_str = deploy_date.replace(".", "")  # MMDD -> MMDD but we need YYYYMMDD
                from datetime import datetime
                year = datetime.now().year
                date_str = f"{year}{deploy_date.replace('.', '')}"
                random_digits = ''.join(random.choices(string.digits, k=3))
                filename = f"SQL_{date_str}_{random_digits}.sql"
            
            # Download the file
            content = self._download_sql_file(cell_l, run_id)
            
            sql_files.append(SQLFile(
                sr_number=sr_number or f"UNKNOWN_{random_digits}",
                content=content,
                filename=filename
            ))
        
        return sql_files

    def _download_sql_file(self, cell_l, run_id: int) -> str:
        """Download SQL file from column L cell."""
        add_event(run_id, "Downloading SQL", "started", "Downloading SQL attachment")
        
        # Navigate to the cell using Name Box
        cell_address = f"L{cell_l.row}"
        self.page.locator('[class*="name-box"]').click()
        self.page.keyboard.type(cell_address)
        self.page.keyboard.press("Enter")
        self.page.wait_for_timeout(500)
        
        # Hover over the cell to show download icon
        cell_selector = f'[data-row="{cell_l.row}"][data-col="11"]'
        self.page.locator(cell_selector).hover()
        self.page.wait_for_timeout(500)
        
        # Click the second icon (download)
        download_icons = self.page.locator('[class*="download-icon"], [title*="下载"], [aria-label*="下载"]')
        if download_icons.count() >= 2:
            download_icons.nth(1).click()
        else:
            # Fallback: try the hyperlink directly
            if cell_l.hyperlink:
                with self.page.expect_download() as download_info:
                    self.page.goto(cell_l.hyperlink.target)
                download = download_info.value
                temp_path = f"workspace/temp_{run_id}.sql"
                download.save_as(temp_path)
                with open(temp_path, 'r', encoding='utf-8', errors='replace') as f:
                    content = f.read()
                Path(temp_path).unlink(missing_ok=True)
                return content
        
        # Wait for download
        with self.page.expect_download() as download_info:
            pass
        download = download_info.value
        temp_path = f"workspace/temp_{run_id}.sql"
        download.save_as(temp_path)
        
        with open(temp_path, 'r', encoding='utf-8', errors='replace') as f:
            content = f.read()
        Path(temp_path).unlink(missing_ok=True)
        
        add_event(run_id, "Downloading SQL", "completed", "SQL file downloaded")
        return content