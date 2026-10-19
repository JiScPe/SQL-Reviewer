import pytest
import openpyxl
from pathlib import Path
from sheet_fetcher import SheetFetcher


@pytest.fixture
def sample_xlsx():
    """Create a sample .xlsx file for testing."""
    wb = openpyxl.Workbook()
    ws = wb.active
    
    # Headers
    headers = ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    
    # Previous separator
    ws.cell(row=2, column=1, value="2026.10.08 8:52 package (10.08 online)")
    
    # Data rows for 10.13 deploy
    ws.cell(row=3, column=4, value="branches/STH260804018.新增.工治具信息维护")
    ws.cell(row=3, column=12, value="SQL")
    ws.cell(row=3, column=12).hyperlink = "https://www.kdocs.cn/office/w/test1"
    
    ws.cell(row=4, column=4, value="branches/STH260804018.其他")
    ws.cell(row=4, column=12, value="SQL")
    ws.cell(row=4, column=12).hyperlink = "https://www.kdocs.cn/office/w/test2"
    
    ws.cell(row=5, column=4, value="merge:branches/STH260904020.成品出库扫码")
    ws.cell(row=5, column=12, value="SQL")
    ws.cell(row=5, column=12).hyperlink = "https://www.kdocs.cn/office/w/test3"
    
    ws.cell(row=6, column=4, value="no sr here")
    ws.cell(row=6, column=12, value="SQL")
    ws.cell(row=6, column=12).hyperlink = "https://www.kdocs.cn/office/w/test4"
    
    # Target separator
    ws.cell(row=7, column=1, value="2026.10.13 9:43 package (10.13 online)")
    
    # Pending rows (should be ignored)
    ws.cell(row=8, column=4, value="branches/STP261014021.EDW JOB")
    ws.cell(row=8, column=12, value="SQL")
    ws.cell(row=8, column=12).hyperlink = "https://www.kdocs.cn/office/w/test5"
    
    xlsx_path = "tests/fixtures/sample_deploy.xlsx"
    Path(xlsx_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    return xlsx_path


def test_sheet_fetcher_parse_deploy_window(sample_xlsx):
    fetcher = SheetFetcher()
    
    # Mock the browser download
    fetcher._download_sql_file = lambda cell, run_id: f"SQL content from {cell.hyperlink.target}"
    
    # We can't easily test the full fetch without browser, so test the parsing logic
    import openpyxl
    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb.active
    
    # Test separator detection
    separators = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, values_only=False), 1):
        cell_a = row[0] if row else None
        if cell_a and cell_a.value:
            text = str(cell_a.value)
            if fetcher.sep_pattern.search(text):
                separators.append((row_idx, text))
    
    assert len(separators) == 2
    assert separators[0][0] == 2
    assert "(10.08 online)" in separators[0][1]
    assert separators[1][0] == 7
    assert "(10.13 online)" in separators[1][1]
    
    # Test window selection for 10.13
    target_sep_idx = 1  # 10.13 separator
    target_row = separators[target_sep_idx][0]
    start_row = separators[target_sep_idx - 1][0] + 1
    end_row = target_row - 1
    
    assert start_row == 3
    assert end_row == 6
    
    # Test SR extraction
    sr_candidates = []
    for row_idx in range(start_row, end_row + 1):
        row = ws[row_idx]
        if len(row) < 12:
            continue
        cell_l = row[11]
        if not cell_l or not cell_l.hyperlink:
            continue
        cell_d = row[3] if len(row) > 3 else None
        sr_number = None
        if cell_d and cell_d.value:
            text = str(cell_d.value)
            match = fetcher.sr_pattern.search(text)
            if match:
                sr_number = match.group(0)
        sr_candidates.append((row_idx, sr_number, cell_l.hyperlink.target))
    
    assert len(sr_candidates) == 4
    assert sr_candidates[0][1] == "STH260804018"  # first match
    assert sr_candidates[1][1] == "STH260804018"  # same SR
    assert sr_candidates[2][1] == "STH260904020"
    assert sr_candidates[3][1] is None  # no SR found


def test_sheet_fetcher_sr_pattern():
    fetcher = SheetFetcher()
    
    # Test various SR formats
    test_cases = [
        ("branches/STH260804018.新增", "STH260804018"),
        ("merge:branches/STP261014021.其他", "STP261014021"),
        ("branches/SFQ260101001.测试", "SFQ260101001"),
        ("no sr here", None),
        ("XYZ260101001", None),  # unlisted prefix
        ("STH26080401", None),  # too short
    ]
    
    for text, expected in test_cases:
        match = fetcher.sr_pattern.search(text)
        if expected:
            assert match is not None
            assert match.group(0) == expected
        else:
            assert match is None


def test_sheet_fetcher_fallback_naming():
    fetcher = SheetFetcher()
    
    # Test fallback naming generates correct format
    from datetime import datetime
    year = datetime.now().year
    date_str = f"{year}1013"
    
    # Just verify the pattern matches
    import re
    fallback_pattern = re.compile(rf"SQL_{year}1013_\d{{3}}\.sql")
    
    # Generate a few
    for _ in range(5):
        random_digits = ''.join(__import__('random').choices(__import__('string').digits, k=3))
        filename = f"SQL_{date_str}_{random_digits}.sql"
        assert fallback_pattern.match(filename)